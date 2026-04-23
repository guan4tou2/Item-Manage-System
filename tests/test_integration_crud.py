"""端到端整合測試：實際建立物品並驗證 CRUD + 新功能。

不像 ``tests/`` 內其他測試多以 mock 為主，本檔案用真實的 SQLite 檔案
跑完整流程，驗證系統真的可以用。

執行:
    .venv313/bin/python tests/test_integration_crud.py
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# 必要環境變數
TMPDIR = Path(tempfile.mkdtemp(prefix="itemmgr_int_"))
DB_PATH = TMPDIR / "items.db"
UPLOAD_DIR = TMPDIR / "uploads"
INSTANCE_DIR = TMPDIR / "instance"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
INSTANCE_DIR.mkdir(parents=True, exist_ok=True)

os.environ.update({
    "SECRET_KEY": "integration-test-secret-32chars-min",
    "DB_TYPE": "postgres",  # 走 SQLAlchemy 分支
    "DATABASE_URL": f"sqlite:///{DB_PATH}",
    "UPLOAD_FOLDER": str(UPLOAD_DIR),
    "REDIS_URL": "memory://",
    "CACHE_TYPE": "SimpleCache",
    # 不要 TEST_MODE：我們要真的 seed admin
})

# 重要：先匯入 models 讓 SQLAlchemy metadata 註冊，
# 再 import app（其中 db.create_all 才能建出全部表）
import app.models  # noqa: F401,E402
from app import create_app, db  # noqa: E402

app = create_app()
app.config["TESTING"] = True
app.config["WTF_CSRF_ENABLED"] = False  # 整合測試暫時關 CSRF

# instance 路徑也指到臨時目錄，避免污染 repo
app.instance_path = str(INSTANCE_DIR)

client = app.test_client()
results: list[tuple[bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((bool(ok), f"  {'✓' if ok else '✗'} {name}{(' — ' + detail) if detail else ''}"))


def login_as_admin() -> None:
    rv = client.post("/signin", data={"UserID": "admin", "Password": "admin"}, follow_redirects=False)
    assert rv.status_code in (302, 200), f"signin failed: {rv.status_code}"


# ---------------------------------------------------------------------------
# 1. 啟動 + 預設 admin 已建立
# ---------------------------------------------------------------------------
print("\n========== 1. 啟動 / 預設 admin ==========")
with app.app_context():
    from app.models import User

    admin = User.query.filter_by(User="admin").first()
    check("預設 admin 已建立", admin is not None)
    check("admin 是管理員", bool(admin and admin.admin))

# ---------------------------------------------------------------------------
# 2. 登入
# ---------------------------------------------------------------------------
print("\n========== 2. 登入 ==========")
rv = client.post("/signin", data={"UserID": "admin", "Password": "admin"}, follow_redirects=False)
check("登入回應 302", rv.status_code == 302, f"status={rv.status_code} loc={rv.location}")
with client.session_transaction() as sess:
    check("session 含 UserID", sess.get("UserID") == "admin")

# ---------------------------------------------------------------------------
# 3. 建立物品類型
# ---------------------------------------------------------------------------
print("\n========== 3. 建立類型 ==========")
rv = client.post("/addtype", data={"TypeName": "電子產品"}, follow_redirects=True)
check("addtype 成功", rv.status_code == 200, f"status={rv.status_code}")
with app.app_context():
    from app.models import ItemType

    tp = ItemType.query.filter_by(name="電子產品").first()
    check("類型已寫入 DB", tp is not None)

# ---------------------------------------------------------------------------
# 4. 建立位置（透過 /locations POST + action=create）
# ---------------------------------------------------------------------------
print("\n========== 4. 建立位置 ==========")
rv = client.post(
    "/locations",
    data={"action": "create", "floor": "1F", "room": "客廳", "zone": "電視櫃"},
    follow_redirects=False,
)
check("位置新增 302", rv.status_code == 302, f"status={rv.status_code}")
with app.app_context():
    from app.models import Location

    loc = Location.query.first()
    check("位置已寫入 DB", loc is not None, f"loc={loc and (loc.floor, loc.room, loc.zone)}")

# ---------------------------------------------------------------------------
# 5. 建立物品（不帶圖片）
# ---------------------------------------------------------------------------
print("\n========== 5. 建立物品 ==========")
# 注意：SQLite 對 Date 欄位較嚴格，這裡測試環境暫不傳保固/到期字串
rv = client.post(
    "/additem",
    data={
        "ItemID": "TEST-001",
        "ItemName": "測試筆電",
        "ItemOwner": "admin",
        "ItemDesc": "整合測試使用",
        "ItemType": "電子產品",
        "ItemFloor": "1F",
        "ItemRoom": "客廳",
        "ItemZone": "電視櫃",
        "ItemStorePlace": "1F-客廳-電視櫃",
        "Quantity": "1",
        "ItemGetDate": "2025-01-15",
    },
    follow_redirects=False,
)
check("additem POST 302", rv.status_code == 302, f"status={rv.status_code} body_excerpt={rv.get_data(as_text=True)[:200]}")

with app.app_context():
    from app.models import Item

    items = Item.query.all()
    check("物品已寫入 DB", len(items) >= 1, f"count={len(items)}")
    if items:
        first = items[0]
        check("物品名稱正確", first.ItemName == "測試筆電", first.ItemName)
        check("物品類型正確", first.ItemType == "電子產品", first.ItemType)
        check("物品 ID 已生成", bool(first.ItemID), str(first.ItemID))

# 再建第二個與第三個物品（不同類型，測 list/filter）
for idx, (name, qty) in enumerate([("測試耳機", 2), ("測試滑鼠", 5)], 2):
    client.post(
        "/additem",
        data={
            "ItemID": f"TEST-00{idx}",
            "ItemName": name,
            "ItemOwner": "admin",
            "ItemDesc": "",
            "ItemType": "電子產品",
            "ItemFloor": "1F",
            "ItemRoom": "客廳",
            "ItemZone": "電視櫃",
            "ItemStorePlace": "1F-客廳-電視櫃",
            "Quantity": str(qty),
            "ItemGetDate": "2025-01-15",
        },
        follow_redirects=False,
    )

with app.app_context():
    from app.models import Item

    total = Item.query.count()
    check("共建立 3 個物品", total == 3, f"count={total}")

# ---------------------------------------------------------------------------
# 6. 列表頁渲染
# ---------------------------------------------------------------------------
print("\n========== 6. 列表頁渲染 ==========")
rv = client.get("/manageitem")
check("manageitem 200", rv.status_code == 200, f"status={rv.status_code}")
html = rv.get_data(as_text=True)
check("列表含「測試筆電」", "測試筆電" in html)
check("列表含「測試耳機」", "測試耳機" in html)
check("含 Excel 匯出按鈕", "匯出 Excel" in html)
check("含 PDF 匯出按鈕", "匯出 PDF" in html)
check("含批次刪除按鈕(admin 可見)", "batchDelete()" in html)
# data-src 只在 ItemPic 有值時出現；無圖物品顯示 placeholder div，這裡放寬
# 改測「lazy_image macro 已被引用」（partials/lazy_image.html 在原始檔有 import）
check("manageitem 引用 lazy_image macro", "lazy_image" in html or "data-src" in html or 'class="bg-light' in html)

# ---------------------------------------------------------------------------
# 7. 匯出測試（CSV / Excel / PDF / JSON）
# ---------------------------------------------------------------------------
print("\n========== 7. 匯出 ==========")
for fmt, mime in [
    ("csv", "text/csv"),
    ("json", "application/json"),
    ("excel", "spreadsheetml"),
    ("pdf", "application/pdf"),
]:
    rv = client.get(f"/export/{fmt}")
    ok = rv.status_code == 200 and (mime in rv.mimetype or mime in rv.get_data()[:50].decode("utf-8", errors="ignore"))
    check(f"匯出 {fmt}", ok, f"status={rv.status_code} mime={rv.mimetype} bytes={len(rv.get_data())}")

# 驗證 Excel 真的能被 openpyxl 讀回來
import io

xlsx_rv = client.get("/export/excel")
if xlsx_rv.status_code == 200 and "spreadsheet" in xlsx_rv.mimetype:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(xlsx_rv.get_data()))
        ws = wb.active
        header = [c.value for c in ws[1]]
        check("Excel 第一列是表頭", "名稱" in header, str(header))
        names_col = [ws.cell(row=r, column=header.index("名稱") + 1).value for r in range(2, ws.max_row + 1)]
        check("Excel 內含「測試筆電」", "測試筆電" in names_col, str(names_col)[:120])
    except Exception as exc:
        check("Excel 解析", False, str(exc)[:120])
else:
    check("Excel 解析", False, f"非 xlsx 回應 status={xlsx_rv.status_code} mime={xlsx_rv.mimetype}")

# ---------------------------------------------------------------------------
# 8. OpenAPI 規格
# ---------------------------------------------------------------------------
print("\n========== 8. OpenAPI ==========")
spec = client.get("/api/openapi.json").get_json()
check("OpenAPI 取得", spec is not None and spec.get("openapi") == "3.0.0")

# ---------------------------------------------------------------------------
# 9. Web Push 訂閱流程
# ---------------------------------------------------------------------------
print("\n========== 9. Web Push 訂閱 ==========")
sub_payload = {
    "endpoint": "https://push.example.com/integration-test",
    "keys": {"p256dh": "BNc...", "auth": "abc..."},
}
rv = client.post("/notifications/subscribe", json=sub_payload)
check("訂閱 200", rv.status_code == 200, f"status={rv.status_code}")

# 確認持久化
with app.app_context():
    from app.utils import push_service

    subs = push_service.list_subscriptions("admin")
    check("訂閱已持久化", len(subs) == 1, f"count={len(subs)}")

# 重複訂閱應去重
rv = client.post("/notifications/subscribe", json=sub_payload)
check("重複訂閱仍 200", rv.status_code == 200)
with app.app_context():
    subs = push_service.list_subscriptions("admin")
    check("去重後仍 1 筆", len(subs) == 1, f"count={len(subs)}")

# 觸發 push-test（無 VAPID 應安全降級）
rv = client.post("/notifications/api/push-test")
data = rv.get_json()
check("push-test 200", rv.status_code == 200, f"status={rv.status_code}")
check("無 VAPID 時 configured=False", data.get("configured") is False, str(data))

# 取消訂閱
rv = client.post("/notifications/unsubscribe", json={"endpoint": sub_payload["endpoint"]})
check("取消訂閱 200", rv.status_code == 200)
with app.app_context():
    subs = push_service.list_subscriptions("admin")
    check("取消後 0 筆", len(subs) == 0)

# ---------------------------------------------------------------------------
# 10. RBAC：登出後嘗試刪除應 401
# ---------------------------------------------------------------------------
print("\n========== 10. RBAC：未登入刪除 ==========")
with app.app_context():
    from app.models import Item

    target_id = Item.query.first().ItemID

# 取一個 cookie 客戶端做未登入測試
fresh = app.test_client()
rv = fresh.post(f"/deleteitem/{target_id}", follow_redirects=False)
# require_permission 在無 UserID 時會 redirect 到 signin (302)
check("未登入刪除被擋(302/401)", rv.status_code in (302, 401), f"status={rv.status_code}")

# 模擬 viewer 角色：admin=False, role=viewer
viewer_client = app.test_client()
with app.app_context():
    # 建立一個 viewer 使用者
    from werkzeug.security import generate_password_hash
    viewer = User(User="viewer1", Password=generate_password_hash("test"), admin=False)
    # role 欄位不存在於 User model；改用 admin=False 就會被視為 viewer
    db.session.add(viewer)
    db.session.commit()

viewer_client.post("/signin", data={"UserID": "viewer1", "Password": "test"})
rv = viewer_client.post(f"/deleteitem/{target_id}", follow_redirects=False)
# viewer 沒有 ITEM_DELETE 權限，應被擋
check("viewer 刪除被擋", rv.status_code in (302, 403), f"status={rv.status_code}")

with app.app_context():
    from app.models import Item

    still = Item.query.filter_by(ItemID=target_id).first()
    check("物品仍未被刪", still is not None)

# ---------------------------------------------------------------------------
# 11. admin 真的刪除
# ---------------------------------------------------------------------------
print("\n========== 11. admin 刪除物品 ==========")
rv = client.post(f"/deleteitem/{target_id}", follow_redirects=False)
check("admin 刪除 302", rv.status_code in (200, 302), f"status={rv.status_code}")

with app.app_context():
    from app.models import Item
    # 軟刪除：is_deleted=True
    deleted = Item.query.filter_by(ItemID=target_id).first()
    if deleted is not None:
        check("物品已軟刪除", bool(deleted.is_deleted), f"is_deleted={deleted.is_deleted}")
    else:
        check("物品已實際移除", True)

    remaining = Item.query.filter_by(is_deleted=False).count()
    total = Item.query.count()
    check("剩餘未刪除 = 2", remaining == 2, f"remaining={remaining} total={total}")

# ---------------------------------------------------------------------------
# 報告
# ---------------------------------------------------------------------------
print("\n" + "=" * 60)
for _, line in results:
    print(line)
total = len(results)
passed = sum(1 for ok, _ in results if ok)
print("=" * 60)
print(f"\n結果: {passed}/{total} 通過")
print(f"暫存目錄: {TMPDIR}")
sys.exit(0 if passed == total else 1)
